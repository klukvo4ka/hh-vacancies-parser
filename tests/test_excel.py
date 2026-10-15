from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.exporters import export_vacancies_to_excel
from tests.helpers import make_vacancy

HEADERS = [
    "Vacancy",
    "Company",
    "Salary From",
    "Salary To",
    "Currency",
    "Salary",
    "Gross",
    "City",
    "Experience",
    "Employment",
    "Work Format",
    "Published At",
    "Link",
]
TEXT_COLUMNS = [1, 2, 6, 8, 9, 10, 11, 12, 13]


def export_and_reopen(vacancies, path: Path) -> Worksheet:
    export_vacancies_to_excel(vacancies, path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Vacancies"]
    return workbook.active


@pytest.fixture
def sheet(tmp_path: Path) -> Worksheet:
    vacancies = [
        make_vacancy(id="1"),
        make_vacancy(
            id="2",
            name="Senior Python",
            salary_from=None,
            salary_to=None,
            salary_currency=None,
            salary_gross=None,
            url="",
        ),
        make_vacancy(id="3", name="Backend", salary_from=0, salary_to=None, salary_currency="USD", salary_gross=True),
    ]
    return export_and_reopen(vacancies, tmp_path / "vacancies.xlsx")


def test_headers_and_row_order(sheet: Worksheet) -> None:
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert [sheet.cell(row, 1).value for row in (2, 3, 4)] == ["Python Developer", "Senior Python", "Backend"]
    assert sheet.max_row == 4


def test_complete_row_values(sheet: Worksheet) -> None:
    assert [cell.value for cell in sheet[2]] == [
        "Python Developer",
        "ACME",
        150000,
        250000,
        "RUR",
        "150 000–250 000 ₽, на руки",
        "Нет",
        "Москва",
        "От 1 года до 3 лет",
        "Полная занятость",
        "Удалённо",
        "2026-10-15T10:30:00+0300",
        "https://hh.ru/vacancy/1",
    ]


def test_salary_cells_stay_numeric_with_grouped_format(sheet: Worksheet) -> None:
    for cell in (sheet["C2"], sheet["D2"]):
        assert cell.data_type == "n" and type(cell.value) is int
        assert cell.number_format == "#,##0"
    assert sheet["C4"].value == 0 and sheet["C4"].data_type == "n"


def test_missing_values_are_blank_cells(sheet: Worksheet) -> None:
    assert sheet["C3"].value is None and sheet["D3"].value is None and sheet["E3"].value is None
    assert sheet["F3"].value == "Не указана"
    assert sheet["G3"].value is None


def test_gross_labels(sheet: Worksheet) -> None:
    assert [sheet.cell(row, 7).value for row in (2, 3, 4)] == ["Нет", None, "Да"]


def test_links_are_native_hyperlinks(sheet: Worksheet) -> None:
    link = sheet["M2"]
    assert link.hyperlink is not None and link.hyperlink.target == "https://hh.ru/vacancy/1"
    assert link.style == "Hyperlink" and link.font.u == "single"
    assert link.data_type == "s"

    empty = sheet["M3"]
    assert empty.value is None and empty.hyperlink is None


def test_header_styling(sheet: Worksheet) -> None:
    for cell in sheet[1]:
        assert cell.font.bold
        assert cell.font.color.rgb.endswith("FFFFFF")
        assert cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith("1F4E78")
        assert cell.alignment.horizontal == "center" and cell.alignment.wrap_text
        assert cell.border.bottom.style == "thin"
    assert sheet.row_dimensions[1].height == 28


def test_sheet_layout(sheet: Worksheet) -> None:
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:M4"
    assert sheet.sheet_view.showGridLines is False
    widths = {letter: sheet.column_dimensions[letter].width for letter in "ABCDEFGHIJKLM"}
    assert widths["A"] == 32 and widths["M"] == max(widths.values())
    assert all(width for width in widths.values())


def test_data_alignment(sheet: Worksheet) -> None:
    for column in TEXT_COLUMNS:
        alignment = sheet.cell(2, column).alignment
        assert alignment.horizontal == "left" and alignment.vertical == "top" and alignment.wrap_text
    assert sheet["C2"].alignment.horizontal == "right"
    assert sheet["E2"].alignment.horizontal == "center" and sheet["G2"].alignment.horizontal == "center"


def test_empty_export_is_styled_header_only(tmp_path: Path) -> None:
    sheet = export_and_reopen([], tmp_path / "empty.xlsx")
    assert sheet.max_row == 1 and [cell.value for cell in sheet[1]] == HEADERS
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:M1"


def test_text_starting_with_equals_is_not_a_formula(tmp_path: Path) -> None:
    sheet = export_and_reopen([make_vacancy(name='=HYPERLINK("http://evil")')], tmp_path / "x.xlsx")
    assert sheet["A2"].value == '=HYPERLINK("http://evil")' and sheet["A2"].data_type == "s"


def test_illegal_control_characters_are_removed(tmp_path: Path) -> None:
    sheet = export_and_reopen([make_vacancy(company="Bad\x0bChars\x00Co")], tmp_path / "x.xlsx")
    assert sheet["B2"].value == "BadCharsCo"


def test_generator_input_and_overwrite(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    export_and_reopen((make_vacancy(id=str(i)) for i in range(5)), path)
    sheet = export_and_reopen(iter([make_vacancy()]), path)
    assert sheet.max_row == 2 and sheet.auto_filter.ref == "A1:M2"


def test_creates_parent_directories_and_returns_path(tmp_path: Path) -> None:
    path = tmp_path / "nested" / "dir" / "vacancies.xlsx"
    assert export_vacancies_to_excel([make_vacancy()], path) == path
    assert path.is_file()


def test_vacancies_are_not_mutated(tmp_path: Path) -> None:
    vacancies = [make_vacancy(), make_vacancy(id="2", url="")]
    snapshot = list(vacancies)
    export_vacancies_to_excel(vacancies, tmp_path / "x.xlsx")
    assert vacancies == snapshot
